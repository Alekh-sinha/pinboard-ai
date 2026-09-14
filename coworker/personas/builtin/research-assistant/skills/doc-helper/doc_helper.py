# /// script
# requires-python = ">=3.10"
# dependencies = [
#   "python-docx>=1.1",
#   "pypdf>=4.0",
#   "pdfplumber>=0.11",
# ]
# ///
"""doc-helper — a cheap first read on a Word/PDF document, and a plain Markdown ->
Word converter for the final deliverable. Run with `uv run doc_helper.py <command> ...`
so its three dependencies resolve in an ephemeral env — never installed into the
caller's own environment.

Two commands:
  skim         <file> [--budget-tokens N]   -> excerpt + heading outline, as JSON
  write-docx   <markdown> --out <path>      -> a .docx built from Markdown headings/
                                                 bullets/bold/italic (not a full Markdown
                                                 implementation — see the docstring below)

Excel is deliberately NOT handled here — a spreadsheet is already structured data;
use the rag MCP server's `query_spreadsheet` instead.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Any

CHARS_PER_TOKEN = 4  # rough, documented estimate — not a real tokenizer


def _docx_text_and_outline(path: Path) -> tuple[str, list[dict]]:
    from docx import Document

    doc = Document(str(path))
    parts: list[str] = []
    outline: list[dict] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        parts.append(text)
        style = (para.style.name if para.style else "") or ""
        if style == "Title":
            outline.append({"level": 0, "text": text})
        elif style.startswith("Heading"):
            digits = "".join(ch for ch in style if ch.isdigit())
            outline.append({"level": int(digits) if digits else 1, "text": text})
    return "\n".join(parts), outline


def _pdf_text(path: Path, budget_chars: int) -> tuple[str, bool]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    out: list[str] = []
    total = 0
    truncated = False
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        if total + len(text) > budget_chars:
            out.append(text[: budget_chars - total])
            truncated = True
            break
        out.append(text)
        total += len(text)
    return "\n".join(out), truncated


_COLOR_SPARSE_MAX_WORDS = 20  # a color used this rarely or less reads as "labels", not body text
_COLOR_LINE_MAX_WORDS = 12  # a heading-by-color candidate line must still be phrase-length


def _pdf_outline(path: Path, max_pages: int = 40) -> list[dict]:
    """Heading heuristic, two signals OR'd together — either is enough to flag a line:

    1. SIZE: a line whose words average > 1.15x the page's median font size (the
       classic case — a bigger title/section font).
    2. COLOR: a short line (<=12 words) whose words share a color that (a) isn't the
       page's dominant (most-used) color — i.e. not the body text color — and (b) is
       itself used sparingly across the whole page (<=20 words total) — i.e. it reads
       as a handful of scattered labels, not a whole paragraph rendered in a tint
       (owner-hit 2026-09-14: a real equity report used a flat font hierarchy but
       colored its section labels gray against black body text — disclaimers/
       footnotes were ALSO gray, so color alone isn't enough; the sparse-usage check
       is what tells a five-word heading apart from a sixty-word gray disclaimer).

    Cheap and imperfect — good enough to point a reader at the right section, not a
    substitute for actually reading the page."""
    import pdfplumber

    outline: list[dict] = []
    with pdfplumber.open(str(path)) as pdf:
        for page_num, page in enumerate(pdf.pages[:max_pages], start=1):
            words = page.extract_words(extra_attrs=["size", "non_stroking_color"])
            if not words:
                continue
            sizes = sorted(w["size"] for w in words)
            median = sizes[len(sizes) // 2]
            size_threshold = median * 1.15

            color_counts: dict[Any, int] = {}
            for w in words:
                color_counts[w["non_stroking_color"]] = (
                    color_counts.get(w["non_stroking_color"], 0) + 1
                )
            dominant_color = max(color_counts, key=color_counts.get)

            # Group words into lines by their vertical position (rounded).
            lines: dict[int, list[dict]] = {}
            for w in words:
                key = round(w["top"])
                lines.setdefault(key, []).append(w)

            for key in sorted(lines):
                line_words = lines[key]
                text = " ".join(w["text"] for w in line_words).strip()
                if not text:
                    continue
                avg_size = sum(w["size"] for w in line_words) / len(line_words)
                by_size = avg_size > size_threshold

                by_color = False
                if not by_size and len(line_words) <= _COLOR_LINE_MAX_WORDS:
                    line_colors = [w["non_stroking_color"] for w in line_words]
                    mode_color = max(set(line_colors), key=line_colors.count)
                    uniform_enough = line_colors.count(mode_color) / len(line_colors) >= 0.6
                    by_color = (
                        uniform_enough
                        and mode_color != dominant_color
                        and color_counts.get(mode_color, 0) <= _COLOR_SPARSE_MAX_WORDS
                    )

                if by_size or by_color:
                    outline.append({"page": page_num, "text": text})

    # A running header/footer (a masthead, page banner) repeats verbatim on nearly
    # every page — a real section heading almost never does. Drop anything that
    # showed up 3+ times as noise rather than structure.
    text_counts: dict[str, int] = {}
    for entry in outline:
        text_counts[entry["text"]] = text_counts.get(entry["text"], 0) + 1
    return [e for e in outline if text_counts[e["text"]] < 3]


def cmd_skim(args: argparse.Namespace) -> None:
    path = Path(args.file)
    if not path.is_file():
        print(json.dumps({"error": f"no such file: {path}"}))
        sys.exit(1)
    budget_chars = args.budget_tokens * CHARS_PER_TOKEN
    suffix = path.suffix.lower()
    if suffix == ".docx":
        text, outline = _docx_text_and_outline(path)
        truncated = len(text) > budget_chars
        excerpt = text[:budget_chars]
    elif suffix == ".pdf":
        excerpt, truncated = _pdf_text(path, budget_chars)
        outline = _pdf_outline(path)
    else:
        print(json.dumps({"error": f"unsupported file type for skim: {suffix} (use rag_ingest_document + query_spreadsheet for Excel)"}))
        sys.exit(1)
    print(
        json.dumps(
            {
                "excerpt": excerpt,
                "excerpt_truncated": truncated,
                "outline": outline,
                "note": (
                    "excerpt is roughly the first N tokens (char-count estimate, not a "
                    "real tokenizer); outline is a heading/structure heuristic, not "
                    "guaranteed complete — judge from both whether this excerpt already "
                    "answers what you need, or where in the document to look instead."
                ),
            }
        )
    )


# -- markdown -> docx (deliberately simple; not a full Markdown implementation) -----

_BOLD_ITALIC_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def _add_runs(paragraph, text: str) -> None:
    for chunk in _BOLD_ITALIC_RE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


_TABLE_SEPARATOR_RE = re.compile(r"^\|?[\s:|-]+\|?$")


def _split_table_row(line: str) -> list[str]:
    # Strip one leading/trailing "|" (optional per the spec) before splitting on "|".
    stripped = line.strip()
    if stripped.startswith("|"):
        stripped = stripped[1:]
    if stripped.endswith("|"):
        stripped = stripped[:-1]
    return [cell.strip() for cell in stripped.split("|")]


def _render_table(doc, table_lines: list[str]) -> None:
    rows = [_split_table_row(line) for line in table_lines]
    # The second line is conventionally the header/body separator (":---", "---", …) —
    # drop it if present rather than rendering it as a data row.
    if len(rows) > 1 and _TABLE_SEPARATOR_RE.match(table_lines[1].strip()):
        header, body = rows[0], rows[2:]
    else:
        header, body = rows[0], rows[1:]
    n_cols = len(header)
    table = doc.add_table(rows=1, cols=n_cols)
    table.style = "Table Grid"  # always present in python-docx's default template
    for cell, text in zip(table.rows[0].cells, header):
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
    for row_cells in body:
        cells = table.add_row().cells
        # A short/long row (malformed markdown) still renders — pad or truncate rather
        # than crashing on a ragged table.
        for i in range(n_cols):
            text = row_cells[i] if i < len(row_cells) else ""
            _add_runs(cells[i].paragraphs[0], text)
    doc.add_paragraph()  # breathing room after the table


def cmd_write_docx(args: argparse.Namespace) -> None:
    from docx import Document

    md_path = Path(args.markdown)
    if not md_path.is_file():
        print(json.dumps({"error": f"no such file: {md_path}"}))
        sys.exit(1)
    doc = Document()
    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].rstrip().strip().startswith("|"):
                table_lines.append(lines[i].rstrip())
                i += 1
            _render_table(doc, table_lines)
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.*)$", line)
        bullet_match = re.match(r"^\s*[-*]\s+(.*)$", line)
        if heading_match:
            level = min(len(heading_match.group(1)), 4)
            doc.add_heading(heading_match.group(2).strip(), level=level)
        elif bullet_match:
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, bullet_match.group(1))
        else:
            p = doc.add_paragraph()
            _add_runs(p, line.strip())
        i += 1
    out_path = Path(args.out)
    doc.save(str(out_path))
    print(json.dumps({"out": str(out_path)}))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)

    p_skim = sub.add_parser("skim", help="excerpt + heading outline for a docx/PDF")
    p_skim.add_argument("file")
    p_skim.add_argument("--budget-tokens", type=int, default=3000)
    p_skim.set_defaults(func=cmd_skim)

    p_write = sub.add_parser("write-docx", help="Markdown -> a .docx")
    p_write.add_argument("markdown")
    p_write.add_argument("--out", required=True)
    p_write.set_defaults(func=cmd_write_docx)

    args = parser.parse_args()
    args.func(args)


if __name__ == "__main__":
    main()
